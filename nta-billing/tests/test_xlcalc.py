"""Built-in formula engine checks. Run: python -m unittest discover tests"""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "engine"))
import xlcalc  # noqa: E402


def calc(cells, other=None):
    """cells: {coord: value} on sheet 'S'; returns {coord: cached value}."""
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "t.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        for k, v in cells.items():
            ws[k] = v
        if other:
            o = wb.create_sheet("1549")
            for k, v in other.items():
                o[k] = v
        wb.save(p)
        xlcalc.recalc(p)
        got = openpyxl.load_workbook(p, data_only=True)["S"]
        return {k: got[k].value for k in cells}


class TestXlcalc(unittest.TestCase):
    def test_project_patterns(self):
        r = calc({"A1": "12.065.1613", "A2": "12.065.1615", "A3": "01.1",
                  "D1": 2, "D2": None, "I1": 10, "I2": 5, "I3": 7,
                  "E1": '=IF(D1="","",D1*0.15)', "E2": '=IF(D2="","",D2*0.15)',
                  "F1": '=SUMIF(A:A,"12.*",I:I)', "F2": "=SUM(I1:I3)/0.85/1.05",
                  "F3": '=IF(SUM(D1,D2)>0,2,0)', "F4": "=-(D1+D2)-I1",
                  "G1": '=IFERROR(INDEX(\'1549\'!D:D,MATCH(A2,\'1549\'!A:A,0)),"לא מוצא")',
                  "G2": "=IFERROR(INDEX('1549'!D:D,MATCH(A3,'1549'!A:A,0)),0)",
                  "G3": "=(I1*(1/100*18))"},
                 other={"A1": "12.065.1615", "D1": 3.5})
        self.assertAlmostEqual(r["E1"], 0.3)
        self.assertIn(r["E2"], ("", None))     # openpyxl reads an empty string back as None
        self.assertEqual(r["F1"], 15)
        self.assertAlmostEqual(r["F2"], 22 / 0.85 / 1.05)
        self.assertEqual(r["F3"], 2)
        self.assertEqual(r["F4"], -12)
        self.assertEqual(r["G1"], 3.5)
        self.assertEqual(r["G2"], 0)
        self.assertAlmostEqual(r["G3"], 1.8)

    def test_operators_and_errors(self):
        r = calc({"A1": "=-2^2", "A2": "=1/0", "A3": '="a"&1.5', "A4": "=ROUND(2.5,0)",
                  "A5": "=ROUND(-1.005,2)", "A6": "=50%", "A7": '="x"*2', "A8": "=1<\"a\""})
        self.assertEqual(r["A1"], 4)
        self.assertEqual(r["A2"], "#DIV/0!")
        self.assertEqual(r["A3"], "a1.5")
        self.assertEqual(r["A4"], 3)
        self.assertEqual(r["A5"], -1.01)
        self.assertEqual(r["A6"], 0.5)
        self.assertEqual(r["A7"], "#VALUE!")
        self.assertIs(r["A8"], True)

    def test_unsupported_raises(self):
        with self.assertRaises(xlcalc.UnsupportedFormula):
            calc({"A1": "=XLOOKUP(1,B:B,C:C)"})


if __name__ == "__main__":
    unittest.main()
