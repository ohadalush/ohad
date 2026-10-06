"""
nta-billing command line. Run from the project folder:

  python cli.py inspect  <raw.xlsx>                      # read + suggest, builds nothing
  python cli.py add      <raw.xlsx> --account 7 [--set CODE=VAL ...] [--force]
  python cli.py query    <raw.xlsx> [--set CODE=VAL ...]  # שאילתא: build files, don't save to account
  python cli.py set      7 1549 CODE=VAL [CODE=auto ...]  # change an order's overrides + rebuild it
  python cli.py meta     7 1549 [--address ..] [--floor ..] [--apt ..] [--name ..]
  python cli.py remove   7 1549
  python cli.py list     7
  python cli.py lines    7 1549                         # non-zero lines of one built order
  python cli.py build    7                              # combined account + order list + פורמט נתע + audit

Override values: a number, or a formula that may use {D:code} {E:code}
{I:code} {L:raw field} placeholders (row numbers are resolved at build time).
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / "engine"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

import openpyxl  # noqa: E402

import config  # noqa: E402
from accumulator import new_job, load_job, save_job, add_order  # noqa: E402
from audit import audit_account, order_totals, order_lines  # noqa: E402
from export_detail_pdf import export_detail_pdf  # noqa: E402
from finalize_order import finalize_order  # noqa: E402
from read_raw import read_raw_quantities  # noqa: E402
from rules import ALL_RAW_FIELDS  # noqa: E402
from suggest import read_elements, suggest  # noqa: E402

SUPPLEMENT_NAMES = {"השלמה", "השלמות"}
SUPPLEMENT_NUMBERS = {"9999", "0000"}


# ---------------------------------------------------------------- helpers
def money(x):
    return f"₪{x:,.2f}"


def parse_sets(items):
    out = {}
    for it in items or []:
        if "=" not in it:
            sys.exit(f"bad --set '{it}', expected CODE=VALUE")
        code, val = it.split("=", 1)
        code, val = code.strip(), val.strip()
        if val.lower() == "auto":
            out[code] = None
            continue
        try:
            out[code] = float(val) if ("." in val or "e" in val.lower()) else int(val)
        except ValueError:
            out[code] = val if val.startswith("=") else "=" + val
    return out


def load_or_new(n):
    p = config.account_json(n)
    return load_job(p) if p.exists() else new_job()


def find_order(job, number):
    for o in job["orders"]:
        if str(o["number"]) == str(number):
            return o
    sys.exit(f"הזמנה {number} לא קיימת בחשבון")


def order_paths(folder, number, name):
    return (folder / f"חשבון {number} - {name}.xlsx",
            folder / f"פירוט כמויות {number} - {name}.pdf")


def build_one(folder, number, name, raw, overrides, meta):
    xlsx, pdf = order_paths(folder, number, name)
    finalize_order(order_name=name, order_number=str(number), raw_values=raw,
                   address=meta.get("כתובת") or "", floor=str(meta.get("קומה") or ""),
                   apartment=str(meta.get("מספר דירה") or ""), out_path=str(xlsx),
                   overrides=overrides)
    try:
        export_detail_pdf(str(xlsx), "פירוט כמויות", str(pdf))
    except FileNotFoundError:            # LibreOffice is only needed for the PDF
        print("⚠ PDF פירוט כמויות לא נוצר: LibreOffice לא מותקן", file=sys.stderr)
        pdf = None
    ws = openpyxl.load_workbook(xlsx, data_only=True)[str(number)]
    tot, vat, inc, isum = order_totals(ws)
    if abs(tot - isum) > 0.01:
        sys.exit(f"✗ ביקורת נכשלה: סה\"כ {tot} ≠ סכום סעיפים {isum}")
    return xlsx, pdf, tot, inc


def read_input(path):
    raw_found, notes, unmatched, meta = read_raw_quantities(path)
    raw = {f: 0 for f in ALL_RAW_FIELDS}
    raw.update(raw_found)
    elements = read_elements(path)
    return raw, notes, unmatched, meta, elements


def show_suggestion(raw, meta, elements, notes, unmatched, sets):
    print(f"הזמנה {meta.get('מספר הזמנה')} - {meta.get('שם משפחה')} | {meta.get('כתובת')} | "
          f"קומה {meta.get('קומה')} | דירה {meta.get('מספר דירה')}")
    print("שדות (לא-אפס):", {k: v for k, v in raw.items() if v})
    print("אלמנטים:")
    for e in elements:
        sh = "תריס אור" if e["light"] else ("תריס" if e["shutter"] else "-")
        print(f"   {e['type']:<12} x{e['count']:g}  {e['unit_area']} מ\"ר/יח'  מסלולים={e['tracks'] or '-'}  {sh}")
    for n in notes:
        print("   הערת שדה:", n)
    ov, warn = suggest(raw, meta, elements)
    for code, v in (sets or {}).items():
        if v is None:
            ov.pop(code, None)
        else:
            ov[code] = v
    print("overrides:", json.dumps(ov, ensure_ascii=False))
    asks = [w for lvl, w in warn if lvl == "ask"]
    for lvl, w in warn:
        print(("⚠ לשאול: " if lvl == "ask" else "ℹ ") + w)
    if unmatched:
        asks.append(f"שדות לא מוכרים בסיכומים: {unmatched}")
        print("⚠ לשאול: שדות לא מוכרים:", unmatched)
    return ov, asks


# --------------------------------------------------------------- commands
def cmd_inspect(a):
    raw, notes, unmatched, meta, elements = read_input(a.file)
    show_suggestion(raw, meta, elements, notes, unmatched, parse_sets(a.set))


def cmd_add(a, query=False):
    raw, notes, unmatched, meta, elements = read_input(a.file)
    ov, asks = show_suggestion(raw, meta, elements, notes, unmatched, parse_sets(a.set))
    if asks and not a.force:
        print("\n✋ לא נבנה - יש שאלות פתוחות. אחרי תשובה: --set CODE=VAL ... ואז --force")
        sys.exit(2)
    number = str(a.number or meta.get("מספר הזמנה"))
    name = a.name or str(meta.get("שם משפחה"))
    folder = (config.OUTPUT / "שאילתות") if query else config.account_dir(a.account)
    folder.mkdir(parents=True, exist_ok=True)
    xlsx, pdf, tot, inc = build_one(folder, number, name, raw, ov, meta)
    print(f"\n✓ {number} - {name}: סה\"כ עלות {money(tot)} | כולל מע\"מ {money(inc)}")
    print(f"   {xlsx}" + (f"\n   {pdf}" if pdf else ""))
    if query:
        print("   (שאילתא - לא נשמר לחשבון)")
        return
    job = load_or_new(a.account)
    add_order(job, name, number, raw, ov)
    job.setdefault("meta", {})[number] = {k: meta.get(k) for k in ("כתובת", "קומה", "מספר דירה")}
    save_job(job, config.account_json(a.account))
    print(f"   נוסף לחשבון {a.account} ({len(job['orders'])} הזמנות)")


def cmd_set(a):
    job = load_job(config.account_json(a.account))
    o = find_order(job, a.order)
    meta = job.get("meta", {}).get(str(o["number"]), {})
    folder = config.account_dir(a.account)
    xlsx, _ = order_paths(folder, o["number"], o["name"])
    before = order_totals(openpyxl.load_workbook(xlsx, data_only=True)[str(o["number"])])[0] if xlsx.exists() else None
    for code, v in parse_sets(a.items).items():
        if v is None:
            o["overrides"].pop(code, None)
        else:
            o["overrides"][code] = v
    xlsx, pdf, tot, inc = build_one(folder, o["number"], o["name"], o["raw"], o["overrides"], meta)
    save_job(job, config.account_json(a.account))
    print(f"✓ {o['number']} - {o['name']}: {money(tot)} (כולל מע\"מ {money(inc)})")
    if before is not None:
        print(f"   שינוי אחרי הנחות: {money(tot - before)}  (לפני: {money(before)})")


def cmd_meta(a):
    job = load_job(config.account_json(a.account))
    o = find_order(job, a.order)
    m = job.setdefault("meta", {}).setdefault(str(o["number"]), {})
    if a.address is not None:
        m["כתובת"] = a.address
    if a.floor is not None:
        m["קומה"] = a.floor
    if a.apt is not None:
        m["מספר דירה"] = a.apt
    if a.name:
        o["name"] = a.name
    save_job(job, config.account_json(a.account))
    build_one(config.account_dir(a.account), o["number"], o["name"], o["raw"], o["overrides"], m)
    print("✓ עודכן ונבנה מחדש:", o["number"], o["name"], m)


def cmd_remove(a):
    job = load_job(config.account_json(a.account))
    n = len(job["orders"])
    job["orders"] = [o for o in job["orders"] if str(o["number"]) != str(a.order)]
    job.get("meta", {}).pop(str(a.order), None)
    save_job(job, config.account_json(a.account))
    print(f"✓ הוסרה ({n} -> {len(job['orders'])})")


def cmd_list(a):
    job = load_job(config.account_json(a.account))
    folder = config.account_dir(a.account)
    total = 0.0
    print(f"| מס' | שם | כתובת | קומה | דירה | סה\"כ לפני מע\"מ |\n|---|---|---|---|---|---|")
    for o in job["orders"]:
        m = job.get("meta", {}).get(str(o["number"]), {})
        xlsx, _ = order_paths(folder, o["number"], o["name"])
        tot = order_totals(openpyxl.load_workbook(xlsx, data_only=True)[str(o["number"])])[0] if xlsx.exists() else 0
        total += tot or 0
        print(f"| {o['number']} | {o['name']} | {m.get('כתובת') or ''} | {m.get('קומה') or ''} | "
              f"{m.get('מספר דירה') or ''} | {money(tot or 0)} |")
    print(f"| **סה\"כ** | | | | | **{money(total)}** |")


def cmd_lines(a):
    job = load_job(config.account_json(a.account))
    o = find_order(job, a.order)
    xlsx, _ = order_paths(config.account_dir(a.account), o["number"], o["name"])
    ws = openpyxl.load_workbook(xlsx, data_only=True)[str(o["number"])]
    for code, (q, t) in order_lines(ws).items():
        print(f"{code:13} {q:>14.4f} {t:>12,.2f}")
    tot, vat, inc, _ = order_totals(ws)
    print(f"סה\"כ {money(tot)} | מע\"מ {money(vat)} | כולל {money(inc)}")
    print("overrides:", json.dumps(o["overrides"], ensure_ascii=False))


def write_order_list(job, path):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "רשימת כתובות"
    ws.sheet_view.rightToLeft = True
    thin = Side(style="thin", color="C9C9C9")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c, h in enumerate(["מס' הזמנה", "שם משפחה", "כתובת", "קומה", "דירה"], 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4A5E")
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    for r, o in enumerate(job["orders"], 2):
        m = job.get("meta", {}).get(str(o["number"]), {})
        for c, v in enumerate([o["number"], o["name"], m.get("כתובת") or "", m.get("קומה") or "",
                               m.get("מספר דירה") or ""], 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="Arial", size=11)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if c in (1, 4, 5) else "right")
    for c, w in zip("ABCDE", [12, 20, 24, 8, 8]):
        ws.column_dimensions[c].width = w
    wb.save(path)


def cmd_build(a):
    from finalize_multi import finalize_multi
    from export_neta import export_neta
    from recalc import recalc_or_raise

    n = a.account
    job = load_job(config.account_json(n))
    # supplement order ("השלמה") always last
    job["orders"].sort(key=lambda o: str(o["name"]) in SUPPLEMENT_NAMES or str(o["number"]) in SUPPLEMENT_NUMBERS)
    save_job(job, config.account_json(n))
    folder = config.account_dir(n)

    combined = folder / f"חשבון {n}.xlsx"
    finalize_multi(job, str(combined))
    ok, report, totals = audit_account(str(combined))
    write_order_list(job, folder / f"רשימת הזמנות - חשבון {n}.xlsx")

    neta = folder / f"חשבון {n} - פורמט נתע.xlsx"
    placed, unplaced, price_changes = export_neta(str(combined), n, str(neta))
    recalc_or_raise(neta, 300)
    neta_total = openpyxl.load_workbook(neta, data_only=True)["ריכוז"]["E8"].value
    acc_total = sum(totals.values())

    print("\n".join(report))
    same = abs(neta_total - acc_total) < 0.01
    print(f"{'✓' if same else '✗'} פורמט נתע: {money(neta_total)} {'=' if same else '≠'} {money(acc_total)}")
    if unplaced:
        print("✗ סעיפים שלא נמצא להם מקום בפורמט נתע:", unplaced)
    for code, (old, new) in price_changes.items():
        print(f"ℹ מחיר {code} בפורמט נתע סונכרן: {old} -> {new}")
    vat = acc_total * config.VAT
    print(f"\nחשבון {n}: {len(job['orders'])} הזמנות | סה\"כ {money(acc_total)} | מע\"מ {money(vat)} | "
          f"כולל {money(acc_total + vat)}")
    for p in (combined, folder / f"רשימת הזמנות - חשבון {n}.xlsx", neta):
        print("  ", p)
    if not (ok and same and not unplaced):
        sys.exit(1)


def main():
    p = argparse.ArgumentParser(description="נת\"ע billing")
    sp = p.add_subparsers(dest="cmd", required=True)

    s = sp.add_parser("inspect"); s.add_argument("file"); s.add_argument("--set", nargs="*")
    for nm in ("add", "query"):
        s = sp.add_parser(nm); s.add_argument("file"); s.add_argument("--set", nargs="*")
        s.add_argument("--force", action="store_true"); s.add_argument("--name"); s.add_argument("--number")
        if nm == "add":
            s.add_argument("--account", required=True)
    s = sp.add_parser("set"); s.add_argument("account"); s.add_argument("order"); s.add_argument("items", nargs="+")
    s = sp.add_parser("meta"); s.add_argument("account"); s.add_argument("order")
    for f in ("--address", "--floor", "--apt", "--name"):
        s.add_argument(f)
    s = sp.add_parser("remove"); s.add_argument("account"); s.add_argument("order")
    s = sp.add_parser("list"); s.add_argument("account")
    s = sp.add_parser("lines"); s.add_argument("account"); s.add_argument("order")
    s = sp.add_parser("build"); s.add_argument("account")

    a = p.parse_args()
    {"inspect": cmd_inspect, "add": cmd_add, "query": lambda x: cmd_add(x, query=True),
     "set": cmd_set, "meta": cmd_meta, "remove": cmd_remove, "list": cmd_list,
     "lines": cmd_lines, "build": cmd_build}[a.cmd](a)


if __name__ == "__main__":
    main()
