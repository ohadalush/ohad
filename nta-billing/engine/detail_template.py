import re
import copy
import openpyxl

from config import TEMPLATE_DETAIL
TEMPLATE_PATH = str(TEMPLATE_DETAIL)
TEMPLATE_SHEET = "גיליון2"
TEMPLATE_ORDER_SHEET_REF = "גיליון1"

_XLOOKUP_RE = re.compile(
    r"_xlfn\.XLOOKUP\(([^,]+),([^,]+),([^,]+),([^()]+)\)"
)


def _convert_formula(formula, new_order_sheet_title):
    f = formula.replace(TEMPLATE_ORDER_SHEET_REF, f"'{new_order_sheet_title}'")

    def repl(m):
        lookup_val, lookup_range, return_range, default = m.groups()
        return f"IFERROR(INDEX({return_range},MATCH({lookup_val},{lookup_range},0)),{default})"

    f = _XLOOKUP_RE.sub(repl, f)
    return f


def build_detail_sheet(wb, order_sheet_title, order_name, order_number,
                        address, floor, apartment):
    src_wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    src_ws = src_wb[TEMPLATE_SHEET]

    ws = wb.create_sheet("פירוט כמויות")
    ws.sheet_view.rightToLeft = True

    for col, dim in src_ws.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[col].width = dim.width
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[r].height = dim.height

    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            new_cell = ws.cell(row=cell.row, column=cell.column)
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                val = _convert_formula(val, order_sheet_title)
            new_cell.value = val
            new_cell.font = copy.copy(cell.font)
            new_cell.border = copy.copy(cell.border)
            new_cell.alignment = copy.copy(cell.alignment)
            new_cell.number_format = cell.number_format
            fg = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
            if fg != "FFFFFF00":
                new_cell.fill = copy.copy(cell.fill)

    for merged_range in src_ws.merged_cells.ranges:
        ws.merge_cells(str(merged_range))

    ws["B109"] = (f"=IFERROR(INDEX('{order_sheet_title}'!D:D,"
                  f"MATCH(A112,'{order_sheet_title}'!A:A,0)),0)"
                  f"+IFERROR(INDEX('{order_sheet_title}'!D:D,"
                  f"MATCH(A113,'{order_sheet_title}'!A:A,0)),0)")
    ws["D112"] = (f"=IFERROR(INDEX('{order_sheet_title}'!D:D,"
                  f"MATCH(A112,'{order_sheet_title}'!A:A,0)),0)")
    ws["D113"] = (f"=IFERROR(INDEX('{order_sheet_title}'!D:D,"
                  f"MATCH(A113,'{order_sheet_title}'!A:A,0)),0)")

    ws["B3"] = f"='{order_sheet_title}'!D1"
    address_str = f"{address}, קומה {floor} דירה {apartment}"
    ws["E3"] = address_str

    return ws
