import openpyxl
from rules import ALL_RAW_FIELDS

_CANON_SET = set(ALL_RAW_FIELDS)
_META_LABELS = {"מספר הזמנה", "שם משפחה", "כתובת", "קומה", "מספר דירה"}


def _normalize_field_name(raw_name):
    """Match a raw field name from an uploaded file to our canonical
    schema (ALL_RAW_FIELDS), tolerating small naming variations such as a
    missing 'שטח ' prefix. Returns (canonical_name_or_None, note)."""
    name = raw_name.strip()
    if name in _CANON_SET:
        return name, None
    candidates = [f"שטח {name}", f'מס\' {name}', f"מספר {name}", f"כמות {name}"]
    for cand in candidates:
        if cand in _CANON_SET:
            return cand, f'"{raw_name}" -> "{cand}"'
    if name.startswith("שטח ") and name[4:] in _CANON_SET:
        return name[4:], f'"{raw_name}" -> "{name[4:]}"'
    return None, None


def read_raw_quantities(path):
    """Returns (raw_values, notes, unmatched, meta) where meta holds any
    order-detail fields found in the file (order number, last name,
    address, floor, apartment) - keyed exactly by their Hebrew labels."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["סיכומים"] if "סיכומים" in wb.sheetnames else wb.active
    values = {}
    notes = []
    unmatched = []
    meta = {}
    for row in ws.iter_rows():
        if len(row) < 2:
            continue
        name, val = row[0].value, row[1].value
        if name is None:
            continue
        name = str(name).strip()
        if name in _META_LABELS:
            meta[name] = val
            continue
        if not isinstance(val, (int, float)):
            continue  # skip header/label rows (non-numeric value)
        canon, note = _normalize_field_name(name)
        if canon is None:
            unmatched.append(name)
            continue
        if note:
            notes.append(note)
        values[canon] = val
    return values, notes, unmatched, meta
