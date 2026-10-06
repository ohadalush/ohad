"""Pure-Python formula engine: recalculate an .xlsx in place without LibreOffice.

openpyxl writes formulas with empty cached values. This module evaluates every
formula (the subset of Excel this project uses, plus a few common extras) and
writes the results into the file's <v> elements, so any reader - openpyxl
data_only=True, file previews, Excel - sees the numbers. Excel still
recalculates on open (fullCalcOnLoad).

An unsupported function or syntax raises UnsupportedFormula, so a result is
never silently wrong; recalc.py then falls back to LibreOffice if installed.

    python engine/xlcalc.py <file.xlsx>
"""
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile
from decimal import Decimal, ROUND_HALF_UP, ROUND_UP, ROUND_DOWN
from xml.sax.saxutils import escape, unescape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


class UnsupportedFormula(Exception):
    pass


class XLError:
    __slots__ = ("code",)

    def __init__(self, code):
        self.code = code

    def __repr__(self):
        return self.code

    def __eq__(self, other):
        return isinstance(other, XLError) and other.code == self.code

    def __hash__(self):
        return hash(self.code)


VALUE, DIV0, REF, NAME, NA, NUM = (XLError(c) for c in
                                   ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!"))
ERROR_CODES = {e.code: e for e in (VALUE, DIV0, REF, NAME, NA, NUM, XLError("#NULL!"))}


class _Empty:
    """A blank cell: 0 in arithmetic, "" in text comparisons."""
    def __repr__(self):
        return "EMPTY"


EMPTY = _Empty()


class _Err(Exception):
    """Internal: carries an XLError up through evaluation."""
    def __init__(self, err):
        self.err = err


# ---------------------------------------------------------------- tokenizer
_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<err>\#(?:NULL!|DIV/0!|VALUE!|REF!|NAME\?|NUM!|N/A))
  | (?P<ref>(?:(?:'(?:[^']|'')+'|[^\s'!:(),"+\-*/^&=<>%{}]+)!)?
             (?:\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?
               |\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}
               |\$?\d+:\$?\d+)(?![\w(]))
  | (?P<num>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
  | (?P<func>[A-Za-z_][A-Za-z0-9_.]*(?=\())
  | (?P<bool>(?:TRUE|FALSE)(?![\w(]))
  | (?P<op><>|<=|>=|[-+*/^&=<>%(),:])
""", re.X)


def tokenize(text):
    pos, out = 0, []
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise UnsupportedFormula(f"cannot parse at {text[pos:pos + 20]!r}")
        kind = m.lastgroup
        if kind != "ws":
            out.append((kind, m.group()))
        pos = m.end()
    return out


# ---------------------------------------------------------------- parser
# AST nodes are tuples: ("num", v) ("str", v) ("bool", v) ("err", e)
# ("ref", sheet, r1, c1, r2, c2, is_range) ("fn", NAME, [args]) ("neg", x)
# ("pct", x) ("bin", op, a, b)

_BIN_PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1,
             "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}


def _parse_ref(text, cur_sheet):
    sheet = cur_sheet
    if "!" in text:
        s, text = text.rsplit("!", 1)
        sheet = s[1:-1].replace("''", "'") if s.startswith("'") else s
    text = text.replace("$", "")
    a, _, b = text.partition(":")
    is_range = bool(b)
    b = b or a

    def part(p):
        m = re.fullmatch(r"([A-Za-z]*)(\d*)", p)
        col = column_index_from_string(m.group(1).upper()) if m.group(1) else None
        row = int(m.group(2)) if m.group(2) else None
        return row, col
    r1, c1 = part(a)
    r2, c2 = part(b)
    return ("ref", sheet, r1, c1, r2, c2, is_range)


class _Parser:
    def __init__(self, tokens, sheet):
        self.t, self.i, self.sheet = tokens, 0, sheet

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self, val=None):
        tok = self.peek()
        if tok[0] is None or (val is not None and tok[1] != val):
            raise UnsupportedFormula(f"expected {val!r}, got {tok[1]!r}")
        self.i += 1
        return tok

    def expr(self, min_prec=0):
        left = self.unary()
        while True:
            kind, val = self.peek()
            if kind != "op" or val not in _BIN_PREC or _BIN_PREC[val] < min_prec:
                return left
            self.i += 1
            prec = _BIN_PREC[val]
            right = self.expr(prec + 1)          # all left-associative, like Excel
            left = ("bin", val, left, right)

    def unary(self):
        kind, val = self.peek()
        if kind == "op" and val in "+-":
            self.i += 1
            x = self.unary()
            return ("neg", x) if val == "-" else x
        return self.postfix()

    def postfix(self):
        x = self.primary()
        while self.peek() == ("op", "%"):
            self.i += 1
            x = ("pct", x)
        return x

    def primary(self):
        kind, val = self.take()
        if kind == "num":
            return ("num", float(val))
        if kind == "str":
            return ("str", val[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("bool", val == "TRUE")
        if kind == "err":
            return ("err", ERROR_CODES.get(val, XLError(val)))
        if kind == "ref":
            return _parse_ref(val, self.sheet)
        if kind == "func":
            name = val.upper()
            if name.startswith("_XLFN."):
                name = name[6:]
            self.take("(")
            args = []
            if self.peek() != ("op", ")"):
                while True:
                    if self.peek() in (("op", ","), ("op", ")")):
                        args.append(("missing",))
                    else:
                        args.append(self.expr())
                    if self.peek() == ("op", ","):
                        self.i += 1
                        continue
                    break
            self.take(")")
            if name not in FUNCS:
                raise UnsupportedFormula(f"function {name} not supported")
            return ("fn", name, args)
        if (kind, val) == ("op", "("):
            x = self.expr()
            self.take(")")
            return x
        raise UnsupportedFormula(f"unexpected {val!r}")


def parse(formula, sheet):
    tokens = tokenize(formula[1:] if formula.startswith("=") else formula)
    p = _Parser(tokens, sheet)
    node = p.expr()
    if p.i != len(tokens):
        raise UnsupportedFormula(f"trailing input {tokens[p.i][1]!r}")
    return node


# ---------------------------------------------------------------- values
class Range:
    __slots__ = ("book", "sheet", "r1", "c1", "r2", "c2")

    def __init__(self, book, sheet, r1, c1, r2, c2):
        self.book, self.sheet = book, sheet
        self.r1, self.c1, self.r2, self.c2 = r1, c1, r2, c2

    @property
    def nrows(self):
        return self.r2 - self.r1 + 1

    @property
    def ncols(self):
        return self.c2 - self.c1 + 1

    def cell(self, i, j):                   # 0-based within the range
        return self.book.value(self.sheet, self.r1 + i, self.c1 + j)

    def coords(self):
        for r in range(self.r1, self.r2 + 1):
            for c in range(self.c1, self.c2 + 1):
                yield r, c

    def values(self):
        for r, c in self.coords():
            yield self.book.value(self.sheet, r, c)


def _num(v):
    """Coerce a scalar for arithmetic (Excel rules)."""
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if v is EMPTY or v is None:
        return 0.0
    if isinstance(v, XLError):
        raise _Err(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            raise _Err(VALUE)
    raise _Err(VALUE)


def _text(v):
    if v is EMPTY or v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return _fmt_num(v)
    if isinstance(v, XLError):
        raise _Err(v)
    return str(v)


def _fmt_num(x):
    if x == int(x) and abs(x) < 1e15:
        return str(int(x))
    return repr(float(f"{x:.15g}"))


def _bool(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    if v is EMPTY or v is None:
        return False
    if isinstance(v, XLError):
        raise _Err(v)
    if isinstance(v, str):
        if v.upper() in ("TRUE", "FALSE"):
            return v.upper() == "TRUE"
    raise _Err(VALUE)


def _type_rank(v):
    if isinstance(v, bool):
        return 2
    if isinstance(v, str):
        return 1
    return 0


def _compare(a, b):
    """-1/0/1 using Excel ordering: numbers < text < booleans; blank adapts."""
    for v in (a, b):
        if isinstance(v, XLError):
            raise _Err(v)
    if a is EMPTY:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is EMPTY:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    ra, rb = _type_rank(a), _type_rank(b)
    if ra != rb:
        return -1 if ra < rb else 1
    if ra == 1:
        a, b = a.lower(), b.lower()
    else:
        a, b = float(a), float(b)
    return (a > b) - (a < b)


def _round(x, digits, mode):
    q = Decimal(1).scaleb(-int(digits))
    d = Decimal(repr(x)).quantize(q, rounding=mode)
    return float(d)


def _wild_regex(pattern):
    out = []
    i = 0
    while i < len(pattern):
        ch = pattern[i]
        if ch == "~" and i + 1 < len(pattern):
            out.append(re.escape(pattern[i + 1]))
            i += 2
            continue
        out.append(".*" if ch == "*" else "." if ch == "?" else re.escape(ch))
        i += 1
    return re.compile("".join(out), re.S | re.I)


def _criteria(crit):
    """Predicate for SUMIF/COUNTIF criteria."""
    if isinstance(crit, XLError):
        raise _Err(crit)
    if not isinstance(crit, str) or crit is EMPTY:
        target = 0.0 if crit is EMPTY else crit
        return lambda v: (not isinstance(v, (str, XLError)) and v is not EMPTY
                          and _compare(v, target) == 0)
    m = re.match(r"(<=|>=|<>|<|>|=)?(.*)", crit, re.S)
    op, rest = m.group(1) or "=", m.group(2)
    try:
        target = float(rest)
        is_num = True
    except ValueError:
        target, is_num = rest, False
    if is_num:
        def num_pred(v):
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                return op == "<>"
            c = _compare(float(v), target)
            return {"=": c == 0, "<>": c != 0, "<": c < 0, ">": c > 0, "<=": c <= 0, ">=": c >= 0}[op]
        return num_pred
    if op in ("=", "<>"):
        if rest == "":
            pred = (lambda v: v is EMPTY or v == "")
        elif any(ch in rest for ch in "*?~"):
            rx = _wild_regex(rest)
            pred = (lambda v: isinstance(v, str) and rx.fullmatch(v) is not None)
        else:
            low = rest.lower()
            pred = (lambda v: isinstance(v, str) and v.lower() == low)
        return pred if op == "=" else (lambda v: not pred(v))

    def text_pred(v):
        if not isinstance(v, str):
            return False
        c = _compare(v, target)
        return {"<": c < 0, ">": c > 0, "<=": c <= 0, ">=": c >= 0}[op]
    return text_pred


# ---------------------------------------------------------------- functions
# Each function receives (ev, args) where args are unevaluated AST nodes, so
# IF/IFERROR can be lazy. ev.scalar(node) / ev.items(node) evaluate.

def _flat_numbers(ev, args):
    for a in args:
        if a[0] == "ref":
            for v in ev.items(a):
                if isinstance(v, XLError):
                    raise _Err(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    yield float(v)
        else:
            yield _num(ev.scalar(a))


def f_sum(ev, args):
    return math.fsum(_flat_numbers(ev, args))


def f_max(ev, args):
    vals = list(_flat_numbers(ev, args))
    return max(vals) if vals else 0.0


def f_min(ev, args):
    vals = list(_flat_numbers(ev, args))
    return min(vals) if vals else 0.0


def f_average(ev, args):
    vals = list(_flat_numbers(ev, args))
    if not vals:
        raise _Err(DIV0)
    return math.fsum(vals) / len(vals)


def f_count(ev, args):
    n = 0
    for a in args:
        if a[0] == "ref":
            n += sum(1 for v in ev.items(a) if isinstance(v, (int, float)) and not isinstance(v, bool))
        else:
            try:
                _num(ev.scalar(a))
                n += 1
            except _Err:
                pass
    return float(n)


def f_counta(ev, args):
    n = 0
    for a in args:
        if a[0] == "ref":
            n += sum(1 for v in ev.items(a) if v is not EMPTY)
        else:
            n += 1
    return float(n)


def f_if(ev, args):
    if not 1 <= len(args) <= 3:
        raise _Err(VALUE)
    if _bool(ev.scalar(args[0])):
        return ev.scalar(args[1]) if len(args) > 1 else True
    if len(args) > 2:
        return ev.scalar(args[2])
    return False


def f_iferror(ev, args):
    try:
        v = ev.scalar(args[0])
    except _Err:
        return ev.scalar(args[1])
    return ev.scalar(args[1]) if isinstance(v, XLError) else v


def f_iserror(ev, args):
    try:
        v = ev.scalar(args[0])
    except _Err:
        return True
    return isinstance(v, XLError)


def f_isblank(ev, args):
    return ev.scalar(args[0], keep_empty=True) is EMPTY


def f_isnumber(ev, args):
    try:
        v = ev.scalar(args[0])
    except _Err:
        return False
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def f_and(ev, args):
    vals = [_bool(v) for a in args for v in (ev.items(a) if a[0] == "ref" else [ev.scalar(a)])
            if v is not EMPTY and not (a[0] == "ref" and isinstance(v, str))]
    if not vals:
        raise _Err(VALUE)
    return all(vals)


def f_or(ev, args):
    vals = [_bool(v) for a in args for v in (ev.items(a) if a[0] == "ref" else [ev.scalar(a)])
            if v is not EMPTY and not (a[0] == "ref" and isinstance(v, str))]
    if not vals:
        raise _Err(VALUE)
    return any(vals)


def f_not(ev, args):
    return not _bool(ev.scalar(args[0]))


def _round_fn(mode):
    def fn(ev, args):
        x = _num(ev.scalar(args[0]))
        d = _num(ev.scalar(args[1])) if len(args) > 1 else 0
        return _round(x, d, mode)
    return fn


def f_abs(ev, args):
    return abs(_num(ev.scalar(args[0])))


def f_int(ev, args):
    return float(math.floor(_num(ev.scalar(args[0]))))


def f_ceiling(ev, args):
    x = _num(ev.scalar(args[0]))
    s = _num(ev.scalar(args[1])) if len(args) > 1 else 1.0
    if s == 0:
        return 0.0
    return math.ceil(round(x / s, 12)) * s


def f_floor(ev, args):
    x = _num(ev.scalar(args[0]))
    s = _num(ev.scalar(args[1])) if len(args) > 1 else 1.0
    if s == 0:
        raise _Err(DIV0)
    return math.floor(round(x / s, 12)) * s


def f_sumif(ev, args):
    rng = ev.range(args[0])
    pred = _criteria(ev.scalar(args[1]))
    srng = ev.range(args[2]) if len(args) > 2 else rng
    total = []
    for i in range(rng.nrows):
        for j in range(rng.ncols):
            if pred(rng.cell(i, j)):
                v = srng.book.value(srng.sheet, srng.r1 + i, srng.c1 + j)
                if isinstance(v, XLError):
                    raise _Err(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    total.append(float(v))
    return math.fsum(total)


def f_countif(ev, args):
    rng = ev.range(args[0])
    pred = _criteria(ev.scalar(args[1]))
    return float(sum(1 for v in rng.values() if pred(v)))


def f_sumproduct(ev, args):
    ranges = [ev.range(a) for a in args]
    shape = (ranges[0].nrows, ranges[0].ncols)
    if any((r.nrows, r.ncols) != shape for r in ranges):
        raise _Err(VALUE)
    total = []
    for i in range(shape[0]):
        for j in range(shape[1]):
            p = 1.0
            for r in ranges:
                v = r.cell(i, j)
                if isinstance(v, XLError):
                    raise _Err(v)
                p *= float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
            total.append(p)
    return math.fsum(total)


def f_match(ev, args):
    look = ev.scalar(args[0])
    rng = ev.range(args[1])
    mtype = _num(ev.scalar(args[2])) if len(args) > 2 and args[2][0] != "missing" else 1.0
    if rng.nrows != 1 and rng.ncols != 1:
        raise _Err(NA)
    vals = list(rng.values())
    if mtype == 0:
        if isinstance(look, str) and any(ch in look for ch in "*?~"):
            rx = _wild_regex(look)
            for k, v in enumerate(vals):
                if isinstance(v, str) and rx.fullmatch(v):
                    return float(k + 1)
        else:
            for k, v in enumerate(vals):
                if v is EMPTY or isinstance(v, XLError):
                    continue
                if _type_rank(v) == _type_rank(look) and _compare(v, look) == 0:
                    return float(k + 1)
        raise _Err(NA)
    best = None                          # approximate match on sorted data
    for k, v in enumerate(vals):
        if v is EMPTY or isinstance(v, XLError) or _type_rank(v) != _type_rank(look):
            continue
        c = _compare(v, look)
        if (mtype > 0 and c <= 0) or (mtype < 0 and c >= 0):
            best = k
        else:
            break
    if best is None:
        raise _Err(NA)
    return float(best + 1)


def f_index(ev, args):
    rng = ev.range(args[0])
    r = int(_num(ev.scalar(args[1]))) if len(args) > 1 and args[1][0] != "missing" else 0
    c = int(_num(ev.scalar(args[2]))) if len(args) > 2 and args[2][0] != "missing" else 0
    if rng.nrows == 1 and len(args) == 2:
        r, c = 1, r
    r, c = r or 1, c or 1
    if not (1 <= r <= rng.nrows and 1 <= c <= rng.ncols):
        raise _Err(REF)
    return rng.cell(r - 1, c - 1)


def f_vlookup(ev, args):
    look = ev.scalar(args[0])
    rng = ev.range(args[1])
    col = int(_num(ev.scalar(args[2])))
    approx = _bool(ev.scalar(args[3])) if len(args) > 3 else True
    if not 1 <= col <= rng.ncols:
        raise _Err(REF)
    first = Range(rng.book, rng.sheet, rng.r1, rng.c1, rng.r2, rng.c1)
    k = _match_in(ev, look, first, 1.0 if approx else 0.0)
    return rng.cell(k - 1, col - 1)


def _match_in(ev, look, rng, mtype):
    """Run f_match on already-evaluated inputs."""
    saved = ev._lit
    ev._lit = {"look": look, "rng": rng}
    try:
        return int(f_match(ev, [("lit", "look"), ("lit", "rng"), ("num", mtype)]))
    finally:
        ev._lit = saved


FUNCS = {
    "SUM": f_sum, "MAX": f_max, "MIN": f_min, "AVERAGE": f_average,
    "COUNT": f_count, "COUNTA": f_counta, "COUNTIF": f_countif,
    "IF": f_if, "IFERROR": f_iferror, "ISERROR": f_iserror,
    "ISBLANK": f_isblank, "ISNUMBER": f_isnumber,
    "AND": f_and, "OR": f_or, "NOT": f_not,
    "ROUND": _round_fn(ROUND_HALF_UP), "ROUNDUP": _round_fn(ROUND_UP),
    "ROUNDDOWN": _round_fn(ROUND_DOWN), "ABS": f_abs, "INT": f_int,
    "CEILING": f_ceiling, "FLOOR": f_floor,
    "SUMIF": f_sumif, "SUMPRODUCT": f_sumproduct,
    "MATCH": f_match, "INDEX": f_index, "VLOOKUP": f_vlookup,
}


# ---------------------------------------------------------------- workbook
class Book:
    def __init__(self, wb):
        self.cells = {}                  # (sheet, row, col) -> constant or formula text
        self.dims = {}
        self.formulas = []               # (sheet, row, col) in file order
        for ws in wb.worksheets:
            self.dims[ws.title] = (ws.max_row, ws.max_column)
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None:
                        continue
                    if isinstance(v, str) and v.startswith("=") and len(v) > 1:
                        self.formulas.append((ws.title, c.row, c.column))
                    elif not isinstance(v, (str, int, float, bool)):
                        if hasattr(v, "text"):        # ArrayFormula / DataTableFormula
                            raise UnsupportedFormula(f"{ws.title}!{c.coordinate}: array formula")
                        v = str(v)
                    self.cells[(ws.title, c.row, c.column)] = v
        self.results = {}
        self._busy = set()
        self._ast = {}

    def value(self, sheet, row, col):
        key = (sheet, row, col)
        if key in self.results:
            return self.results[key]
        v = self.cells.get(key, EMPTY)
        if isinstance(v, str) and v.startswith("=") and len(v) > 1:
            return self._eval_cell(key, v)
        if isinstance(v, int) and not isinstance(v, bool):
            return float(v)
        return v

    def _eval_cell(self, key, formula):
        if key in self._busy:
            raise UnsupportedFormula(f"circular reference at {key[0]}!{get_column_letter(key[2])}{key[1]}")
        self._busy.add(key)
        try:
            ast = self._ast.get(key)
            if ast is None:
                try:
                    ast = parse(formula, key[0])
                except UnsupportedFormula as e:
                    raise UnsupportedFormula(f"{key[0]}!{get_column_letter(key[2])}{key[1]} {formula}: {e}")
            ev = _Eval(self)
            try:
                v = ev.scalar(ast)
            except _Err as e:
                v = e.err
            if v is EMPTY:
                v = 0.0
        finally:
            self._busy.discard(key)
        self.results[key] = v
        return v

    def calculate(self):
        old = sys.getrecursionlimit()
        sys.setrecursionlimit(max(old, 20000))   # long dependency chains recurse
        try:
            for key in self.formulas:
                self.value(*key)
        finally:
            sys.setrecursionlimit(old)
        return self.results


class _Eval:
    def __init__(self, book):
        self.book = book
        self._lit = None

    def range(self, node):
        if node[0] == "lit":
            return self._lit[node[1]]
        if node[0] != "ref":
            raise _Err(VALUE)
        _, sheet, r1, c1, r2, c2, _ = node
        if sheet not in self.book.dims:
            raise _Err(REF)
        max_r, max_c = self.book.dims[sheet]
        if r1 is None:                   # whole column(s): clip to the used area
            r1, r2 = 1, max(max_r, 1)
        if c1 is None:                   # whole row(s)
            c1, c2 = 1, max(max_c, 1)
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
        return Range(self.book, sheet, r1, c1, r2, c2)

    def items(self, node):
        return self.range(node).values()

    def scalar(self, node, keep_empty=False):
        kind = node[0]
        if kind == "num" or kind == "str" or kind == "bool":
            return node[1]
        if kind == "lit":
            return self._lit[node[1]]
        if kind == "err":
            raise _Err(node[1])
        if kind == "missing":
            return EMPTY if keep_empty else 0.0
        if kind == "ref":
            rng = self.range(node)
            if rng.nrows == 1 and rng.ncols == 1:
                v = rng.cell(0, 0)
                if isinstance(v, XLError):
                    raise _Err(v)
                return v
            raise _Err(VALUE)            # implicit intersection not supported
        if kind == "fn":
            v = FUNCS[node[1]](self, node[2])
            if isinstance(v, XLError):
                raise _Err(v)
            return v
        if kind == "neg":
            return -_num(self.scalar(node[1]))
        if kind == "pct":
            return _num(self.scalar(node[1])) / 100.0
        if kind == "bin":
            op = node[1]
            a = self.scalar(node[2])
            b = self.scalar(node[3])
            if op == "&":
                return _text(a) + _text(b)
            if op in ("=", "<>", "<", ">", "<=", ">="):
                c = _compare(a, b)
                return {"=": c == 0, "<>": c != 0, "<": c < 0, ">": c > 0,
                        "<=": c <= 0, ">=": c >= 0}[op]
            x, y = _num(a), _num(b)
            if op == "+":
                return x + y
            if op == "-":
                return x - y
            if op == "*":
                return x * y
            if op == "/":
                if y == 0:
                    raise _Err(DIV0)
                return x / y
            if op == "^":
                try:
                    r = x ** y
                except (OverflowError, ZeroDivisionError):
                    raise _Err(NUM)
                if isinstance(r, complex):
                    raise _Err(NUM)
                return r
        raise UnsupportedFormula(f"node {kind}")


# ---------------------------------------------------------------- write back
_CELL = re.compile(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', re.S)
_ATTR_R = re.compile(r'\br="([A-Z]+)(\d+)"')
_ATTR_T = re.compile(r'\s+t="[^"]*"')
_F = re.compile(r'<f\b[^>]*?(?:/>|>.*?</f>)', re.S)


def _xml_value(v):
    """(t attribute or None, <v> text)."""
    if isinstance(v, bool):
        return "b", "1" if v else "0"
    if isinstance(v, XLError):
        return "e", v.code
    if isinstance(v, str):
        return "str", escape(v)
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return "e", NUM.code
        return None, _fmt_num(v) if v == int(v) and abs(v) < 1e15 else repr(v)
    return None, str(v)


def _sheet_paths(z):
    """sheet title -> xml path inside the zip."""
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_to_target = {}
    for m in re.finditer(r'<Relationship\b[^>]*>', rels):
        tag = m.group()
        rid = re.search(r'\bId="([^"]+)"', tag).group(1)
        target = re.search(r'\bTarget="([^"]+)"', tag).group(1)
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        rid_to_target[rid] = target
    out = {}
    for m in re.finditer(r'<(?:\w+:)?sheet\b[^>]*>', wbxml):
        tag = m.group()
        name = re.search(r'\bname="([^"]*)"', tag).group(1)
        rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
        out[unescape(name, {"&quot;": '"', "&apos;": "'"})] = rid_to_target[rid]
    return out


def _patch_sheet(xml, values):
    """values: {(row, col): value} for this sheet's formula cells."""
    def repl(m):
        attrs, inner = m.group(1), m.group(2)
        if inner is None or "<f" not in inner:
            return m.group(0)
        rm = _ATTR_R.search(attrs)
        if not rm:
            return m.group(0)
        key = (int(rm.group(2)), column_index_from_string(rm.group(1)))
        if key not in values:
            return m.group(0)
        t, text = _xml_value(values[key])
        attrs = _ATTR_T.sub("", attrs)
        if t:
            attrs += f' t="{t}"'
        f = _F.search(inner).group(0)
        return f"<c{attrs}>{f}<v>{text}</v></c>"
    return _CELL.sub(repl, xml)


def write_values(path, results):
    by_sheet = {}
    for (sheet, r, c), v in results.items():
        by_sheet.setdefault(sheet, {})[(r, c)] = v
    with zipfile.ZipFile(path) as z:
        paths = _sheet_paths(z)
        targets = {paths[s]: vals for s, vals in by_sheet.items() if s in paths}
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(path)))
        os.close(fd)
        try:
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
                for info in z.infolist():
                    data = z.read(info.filename)
                    if info.filename in targets:
                        data = _patch_sheet(data.decode("utf-8"), targets[info.filename]).encode("utf-8")
                    elif info.filename == "xl/workbook.xml":
                        data = _force_full_calc(data.decode("utf-8")).encode("utf-8")
                    out.writestr(info, data)
        except BaseException:
            os.remove(tmp)
            raise
    shutil.move(tmp, path)


def _force_full_calc(xml):
    """Ask Excel to recompute on open anyway (our values are a cache)."""
    if "<calcPr" not in xml:
        return xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
    if "fullCalcOnLoad" in xml:
        return xml
    return re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', xml, count=1)


# ---------------------------------------------------------------- entry
def recalc(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    book = Book(wb)
    results = book.calculate()
    write_values(path, results)
    errs = [f"{s}!{get_column_letter(c)}{r}={v.code}"
            for (s, r, c), v in results.items() if isinstance(v, XLError)]
    return {"status": "success" if not errs else "errors_found", "engine": "python",
            "formulas": len(results), "total_errors": len(errs), "errors": errs[:50]}


if __name__ == "__main__":
    print(json.dumps(recalc(sys.argv[1]), ensure_ascii=False, indent=2))
