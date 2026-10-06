"""Build the combined account file end-to-end: generate_multi -> recalc ->
hide zero rows (order sheets by D, ריכוז by 'סה"כ כמות חשבון'=0) -> recalc."""
import openpyxl
from recalc import recalc_or_raise
from generate_multi import generate_multi
from apply_hide import build_row_layout, hide_zero_rows

def _recalc(path):
    return recalc_or_raise(path, 300)

def finalize_multi(job, out_path):
    generate_multi(job, out_path)
    _recalc(out_path)
    wb = openpyxl.load_workbook(out_path)
    wbv = openpyxl.load_workbook(out_path, data_only=True)
    code_to_row, header_children = build_row_layout()
    for order in job["orders"]:
        t = str(order["number"])
        d = {}
        for code, row in code_to_row.items():
            v = wbv[t].cell(row=row, column=4).value
            d[code] = v if isinstance(v, (int, float)) else 0
        hide_zero_rows(wb[t], code_to_row, header_children, d)
    rk, rkv = wb["ריכוז"], wbv["ריכוז"]
    qcol = next(c for c in range(1, rk.max_column + 1) if rkv.cell(row=4, column=c).value == 'סה"כ כמות חשבון')
    for r in range(5, rk.max_row + 1):
        v = rkv.cell(row=r, column=qcol).value
        if not isinstance(v, (int, float)) or v == 0:
            rk.row_dimensions[r].hidden = True
    wb.save(out_path)
    return _recalc(out_path)
