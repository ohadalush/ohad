"""
Export a finished multi-order account (Claude format, e.g. 'חשבון 6.xlsx')
into the client's 'פורמט נתע' workbook.

Rules (taught by the user on 2026-10-04, from חשבון 5):
 1. Every code that exists in the נתע sheet '12 - אלומיניום' -> quantity goes
    there, in the 'כמות מדווחת מאוכלס' column of this account's block
    (חשבון N -> its 4-column block; quantity column = 3rd of the block).
 2. Every other code -> sheet 'סעיפים חסרים', column R (block header P1
    is relabelled to 'חשבון N').
 3. Other sheets (חשמל, מיזוג, יתר סעיפים) are ignored.
 4. Only ONE account per file - all other account blocks are cleared.
 5. 'ריכוז' = aluminium sheet total (this account's ₪ column) + סעיפים חסרים total.
 6. Totals must equal the Claude-format total 1:1 (same quantity source,
    same price list) - verify after export.
Quantities are written as plain values (no external XLOOKUP links).
"""
import copy

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

from config import TEMPLATE_NETA
TEMPLATE = str(TEMPLATE_NETA)
ALU = "12 - אלומיניום"
MISSING = "סעיפים חסרים"
RIKUZ = "ריכוז"


def read_account_quantities(claude_xlsx, qty_col=None):
    """code -> total quantity, from the combined ריכוז sheet of the Claude-format
    account ('סה"כ כמות חשבון' column)."""
    wb = openpyxl.load_workbook(claude_xlsx, data_only=True)
    ws = wb["ריכוז"]
    if qty_col is None:
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=4, column=c).value == 'סה"כ כמות חשבון':
                qty_col = c
    qty = {}
    for r in range(5, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not a or str(a).count(".") < 2:
            continue
        v = ws.cell(row=r, column=qty_col).value
        if isinstance(v, (int, float)) and v != 0:
            code = str(a).strip()
            if code in qty:
                raise ValueError(f"duplicate code in ריכוז: {code}")
            qty[code] = v
    return qty


def export_neta(claude_xlsx, account_no, out_path, template=TEMPLATE):
    qty = read_account_quantities(claude_xlsx)

    wb = openpyxl.load_workbook(template, keep_links=False)
    alu = wb[ALU]
    mis = wb[MISSING]

    # ---- locate this account's block in the aluminium sheet (row 3 headers)
    block_start = None
    for c in range(1, alu.max_column + 1):
        if alu.cell(row=3, column=c).value == f"חשבון {account_no}":
            block_start = c
    if block_start is None:
        raise ValueError(f"'חשבון {account_no}' block not found in {ALU}")
    q_col = block_start + 2                # כמות מדווחת מאוכלס
    t_col = block_start + 3                # סה"כ ₪ מאוכלס
    t_letter = get_column_letter(t_col)

    # first account block starts at K; each block = 4 columns (K..BF)
    first, last = column_index_from_string("K"), column_index_from_string("BF")
    alu_rows = {}
    for r in range(5, alu.max_row + 1):
        a = alu.cell(row=r, column=1).value
        if not a:
            continue
        alu_rows[str(a).strip()] = r
        # clear every account's quantity cells (2 qty cols per block),
        # leaving the ₪ = qty*price formulas in place
        for bs in range(first, last + 1, 4):
            for qc in (bs, bs + 2):
                alu.cell(row=r, column=qc).value = None

    # ---- סעיפים חסרים: single block P..S, relabel to this account
    mis["P1"] = f"חשבון {account_no}"
    mis_rows = {}
    for r in range(3, mis.max_row + 1):
        a = mis.cell(row=r, column=1).value
        if not a or str(a).count(".") < 1:
            continue
        mis_rows[str(a).strip()] = r
        for col in ("K", "M", "P", "R"):
            mis[f"{col}{r}"].value = None
        mis[f"R{r}"].value = 0

    # Claude-format price list is the source of truth (user confirmed
    # 2026-10-04: 60.020.0067 = 3,400, not the 3,500 in the נתע sheet).
    # For every code used in this account, the נתע price cell is synced to
    # ours so the totals are 1:1. Changes are returned for reporting.
    from order_template import CANONICAL_ROWS
    my_price = {c: p for c, d, u, p in CANONICAL_ROWS if c and p != ""}
    price_changes = {}

    # Codes the נתע file has no row for (e.g. new air-conditioning items) are
    # appended to סעיפים חסרים with our description/unit/price; the total row
    # moves down if the free rows run out.
    total_row = next(r for r in range(3, mis.max_row + 1)
                     if str(mis[f"I{r}"].value or "").startswith("=SUM(I3:"))
    my_row = {c: (d, u, p) for c, d, u, p in CANONICAL_ROWS if c}
    extra = [c for c in qty if c not in alu_rows and c not in mis_rows and c in my_row]
    if extra:
        mis[f"I{total_row}"] = None          # rewritten below, possibly lower down
        style_row = max(mis_rows.values())
        r = style_row + 1
        for code in extra:
            d, u, p = my_row[code]
            for col in range(1, mis.max_column + 1):
                src = mis.cell(row=style_row, column=col)
                if src.has_style:
                    mis.cell(row=r, column=col)._style = copy.copy(src._style)
            mis[f"A{r}"], mis[f"B{r}"], mis[f"C{r}"], mis[f"D{r}"] = code, d, u, p
            mis[f"E{r}"] = f"=D{r}*(1-$E$1)"
            mis[f"F{r}"] = f"=E{r}*(1+$F$1)"
            mis[f"H{r}"] = f"=M{r}+R{r}"
            mis[f"I{r}"] = f"=N{r}+S{r}"
            mis[f"S{r}"] = f"=R{r}*F{r}"
            mis_rows[code] = r
            r += 1
        total_row = max(total_row, r + 1)
        mis[f"I{total_row}"] = f"=SUM(I3:I{total_row - 1})"

    placed, unplaced = {}, {}
    for code, q in qty.items():
        if code in alu_rows:
            r = alu_rows[code]
            alu.cell(row=r, column=q_col).value = q
            placed[code] = (ALU, q)
            price_cell = alu.cell(row=r, column=4)
        elif code in mis_rows:
            r = mis_rows[code]
            mis[f"R{r}"].value = q
            placed[code] = (MISSING, q)
            price_cell = mis[f"D{r}"]
        else:
            unplaced[code] = q
            continue
        if code in my_price and abs(float(price_cell.value) - float(my_price[code])) > 1e-9:
            price_changes[code] = (price_cell.value, my_price[code])
            price_cell.value = my_price[code]

    # ---- ריכוז
    rk = wb[RIKUZ]
    rk["E2"] = f"=SUM('{ALU}'!{t_letter}:{t_letter})"
    rk["E6"] = f"='{MISSING}'!I{total_row}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return placed, unplaced, price_changes
