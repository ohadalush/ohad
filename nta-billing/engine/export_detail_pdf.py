import openpyxl, subprocess, sys, os
from config import find_soffice

def export_detail_pdf(order_xlsx_path, sheet_title, out_pdf_path):
    """Extract the 'פירוט כמויות' sheet from a finalized order workbook into
    a standalone xlsx (with page setup for a clean printout), then convert
    it to PDF via headless LibreOffice."""
    soffice = find_soffice()  # FileNotFoundError before any temp file is written
    wb = openpyxl.load_workbook(order_xlsx_path, data_only=True)  # need computed values
    src = wb["פירוט כמויות"]

    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    dst = wb2.create_sheet("פירוט כמויות")
    dst.sheet_view.rightToLeft = True

    for col, dim in src.column_dimensions.items():
        if dim.width:
            dst.column_dimensions[col].width = dim.width
    for r, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[r].height = dim.height
    for merged_range in src.merged_cells.ranges:
        dst.merge_cells(str(merged_range))

    import copy
    for row in src.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            nc = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            nc.font = copy.copy(cell.font)
            nc.border = copy.copy(cell.border)
            nc.alignment = copy.copy(cell.alignment)
            nc.number_format = cell.number_format
            if cell.fill:
                nc.fill = copy.copy(cell.fill)

    # page setup: portrait, fit to 1 page wide
    dst.page_setup.orientation = 'portrait'
    dst.page_setup.fitToWidth = 1
    dst.page_setup.fitToHeight = 0
    dst.sheet_properties.pageSetUpPr.fitToPage = True
    dst.print_area = f"A1:E{dst.max_row}"

    tmp_xlsx = out_pdf_path.replace(".pdf", "_tmp.xlsx")
    wb2.save(tmp_xlsx)

    outdir = os.path.dirname(out_pdf_path) or "."
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, tmp_xlsx],
        check=True, capture_output=True, timeout=60
    )
    produced = os.path.join(outdir, os.path.basename(tmp_xlsx).replace(".xlsx", ".pdf"))
    os.replace(produced, out_pdf_path)
    os.remove(tmp_xlsx)
    return out_pdf_path

if __name__ == "__main__":
    export_detail_pdf(sys.argv[1], "פירוט כמויות", sys.argv[2])
