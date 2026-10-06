"""The 4 audit checks + order totals. Works on recalculated files only."""
import openpyxl

from generate import CHAPTERS
from order_template import CANONICAL_ROWS


def _is_item(code):
    return bool(code) and str(code).count(".") == 2


def order_totals(ws):
    """(total, vat, total_incl, item_sum) of one order sheet."""
    lab = next((c for row in ws.iter_rows() for c in row if c.value == 'סה"כ עלות'), None)
    if lab is None:
        raise ValueError(f"no 'סה\"כ עלות' in sheet {ws.title}")
    tot = ws.cell(row=lab.row, column=9).value
    vat = ws.cell(row=lab.row + 1, column=9).value
    inc = ws.cell(row=lab.row + 2, column=9).value
    isum = 0.0
    for r in range(5, lab.row):
        a, i = ws.cell(row=r, column=1).value, ws.cell(row=r, column=9).value
        if _is_item(a) and isinstance(i, (int, float)):
            isum += i
    return tot, vat, inc, isum


def order_lines(ws):
    """code -> (qty, total) for every non-zero item row."""
    lab = next(c for row in ws.iter_rows() for c in row if c.value == 'סה"כ עלות')
    out = {}
    for r in range(5, lab.row):
        a, d, i = ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value, ws.cell(row=r, column=9).value
        if _is_item(a) and isinstance(d, (int, float)) and d:
            out[str(a)] = (d, i or 0)
    return out


def chapter_coverage():
    prefixes = {p for p, _ in CHAPTERS}
    return sorted({c for c, *_ in CANONICAL_ROWS if c and "." in c and c.split(".")[0] not in prefixes})


def audit_account(path):
    """Returns (ok, report_lines, order_totals_dict)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    lines, ok = [], True
    totals = {}
    for t in wb.sheetnames:
        if t in ("ריכוז", "רשימת כתובות"):
            continue
        tot, _, _, isum = order_totals(wb[t])
        totals[t] = tot
        good = tot is not None and abs(tot - isum) < 0.01
        ok &= good
        if not good:
            lines.append(f"✗ בדיקה 1: הזמנה {t}: סה\"כ {tot} ≠ סכום סעיפים {isum}")
    lines.append(f"{'✓' if ok else '✗'} בדיקה 1: סכום סעיפים = סה\"כ עלות בכל {len(totals)} ההזמנות")

    rk = wb["ריכוז"]
    hdr = {rk.cell(row=4, column=c).value: c for c in range(1, rk.max_column + 1)}
    order_cols = {str(rk.cell(row=3, column=c).value): c for c in range(1, rk.max_column + 1)
                  if rk.cell(row=4, column=c).value == 'סה"כ לדירה'}
    ok2 = True
    for t, col in order_cols.items():
        s = sum(rk.cell(row=r, column=col).value for r in range(5, rk.max_row + 1)
                if _is_item(rk.cell(row=r, column=1).value) and isinstance(rk.cell(row=r, column=col).value, (int, float)))
        if t in totals and abs(s - totals[t]) > 0.01:
            ok2 = False
            lines.append(f"✗ בדיקה 2: ריכוז הזמנה {t}: {s} ≠ {totals[t]}")
    lines.append(f"{'✓' if ok2 else '✗'} בדיקה 2: כל עמודת הזמנה בריכוז = הסה\"כ שלה")
    acol = hdr.get('סה"כ סכום חשבון')
    acc = sum(rk.cell(row=r, column=acol).value for r in range(5, rk.max_row + 1)
              if _is_item(rk.cell(row=r, column=1).value) and isinstance(rk.cell(row=r, column=acol).value, (int, float)))
    ok3 = abs(acc - sum(totals.values())) < 0.01
    lines.append(f"{'✓' if ok3 else '✗'} בדיקה 3: סה\"כ חשבון {acc:,.2f} = סכום ההזמנות {sum(totals.values()):,.2f}")
    missing = chapter_coverage()
    ok4 = not missing
    lines.append(f"{'✓' if ok4 else '✗'} בדיקה 4: כיסוי פרקים" + ("" if ok4 else f" - חסרות קידומות: {missing}"))
    return ok and ok2 and ok3 and ok4, lines, totals
