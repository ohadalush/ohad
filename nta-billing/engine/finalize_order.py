"""Unified single-order finalizer: generate (order sheet + detail sheet),
then hide zero-quantity rows/chapters, consistently, every time."""
import openpyxl
from recalc import recalc_or_raise
from generate import build_order_sheet, CHAPTERS
from detail_template import build_detail_sheet
from apply_hide import build_row_layout, hide_zero_rows


def generate_single_with_detail(order_name, order_number, raw_values,
                                 address, floor, apartment, out_path, overrides=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_title = str(order_number)
    ws_order, code_to_row, last_row, always_show = build_order_sheet(
        wb, sheet_title, order_name, order_number, raw_values, overrides)
    build_detail_sheet(wb, sheet_title, order_name, order_number, address, floor, apartment)
    wb.save(out_path)
    return code_to_row, last_row, always_show


def finalize_order(order_name, order_number, raw_values, address, floor, apartment,
                    out_path, overrides=None, tmp_path=None):
    tmp_path = tmp_path or (out_path + ".tmp.xlsx")
    generate_single_with_detail(order_name, order_number, raw_values, address, floor,
                                 apartment, tmp_path, overrides)

    recalc_or_raise(tmp_path)

    wb = openpyxl.load_workbook(tmp_path, data_only=False)
    wbvals = openpyxl.load_workbook(tmp_path, data_only=True)
    sheet_title = str(order_number)
    ws = wb[sheet_title]
    ws_vals = wbvals[sheet_title]

    code_to_row, header_children = build_row_layout()
    d_values = {}
    for code, row in code_to_row.items():
        v = ws_vals.cell(row=row, column=4).value
        d_values[code] = v if isinstance(v, (int, float)) else 0

    hide_zero_rows(ws, code_to_row, header_children, d_values)

    last_item_row = max(code_to_row.values())
    chapter_row_start = last_item_row + 2
    for i, (prefix, label) in enumerate(CHAPTERS):
        r = chapter_row_start + i
        val = ws_vals.cell(row=r, column=9).value or 0
        if val == 0:
            ws.row_dimensions[r].hidden = True

    wb.save(out_path)
    recalc_or_raise(out_path)          # openpyxl save drops cached values - recompute
    import os
    os.remove(tmp_path)
    return out_path
