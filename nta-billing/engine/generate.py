import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from order_template import CANONICAL_ROWS
from rules import RULES, ALL_RAW_FIELDS
from styles import apply_cell_style, set_column_widths, set_row_height

HEADER_FILL = PatternFill("solid", fgColor="C8C8C8")
BOLD = Font(bold=True)
THIN = Side(style="thin")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

CHAPTERS = [
    ("01", "01 - עבודות עפר"),
    ("05", "05 - עבודות איטום"),
    ("09", "09 - עבודות טיח"),
    ("11", "11 - עבודות צביעה"),
    ("12", "12 - עבודות אלומיניום"),
    ("15", "15 - מתקני מיזוג אוויר"),
    ("22", "22 - אלמנטים מתועשים"),
    ("30", "וילונות - 30"),
    ("60", "60 - מחירי שעות עבודה ושכירת ציוד"),
    ("85", "85 - חומרי איטום ובידוד תרמי"),
]


def build_order_sheet(wb, sheet_title, order_name, order_number, raw_values, overrides=None):
    overrides = overrides or {}
    ws = wb.create_sheet(sheet_title)
    ws.sheet_view.rightToLeft = True
    set_column_widths(ws, "ABCDEFGHI")
    ws.column_dimensions["K"].width = 32
    ws.column_dimensions["L"].width = 12

    ws["C1"] = order_name
    ws["C1"].font = BOLD
    ws["D1"] = order_number
    set_row_height(ws, 1)

    headers = ["סעיף", "תאור", "יח'", "כמות", "מחיר", "הנחת קבלן 15%",
               'תוספת דירה מאוכלסת 5%', 'סה"כ מחיר יחידה', 'סה"כ אחרי הנחה']
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=c, value=h)
        apply_cell_style(cell, "header_row", c)
    set_row_height(ws, 6)

    code_to_row = {}
    r = 7
    for code, desc, unit, price in CANONICAL_ROWS:
        if code is None:
            r += 1
            continue
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=5, value=price if price != "" else None)
        if price != "":
            ws.cell(row=r, column=6, value=f"=IF(E{r}=\"\",\"\",E{r}*0.15)")
            ws.cell(row=r, column=7, value=f"=IF(E{r}=\"\",\"\",E{r}*0.05)")
            ws.cell(row=r, column=8, value=f"=IF(E{r}=\"\",\"\",E{r}*1.05*0.85)")
            ws.cell(row=r, column=9, value=f"=IF(D{r}=\"\",\"\",H{r}*D{r})")
        for c in range(1, 10):
            apply_cell_style(ws.cell(row=r, column=c), "item_row", c)
        set_row_height(ws, r)
        if code == "11.011.2000":
            ws.row_dimensions[r].height = 28.8  # not in the original template; wraps to 2 lines
        code_to_row[code] = r
        r += 1
    last_item_row = r - 1

    # ---- chapter summary block ----
    chapter_total_cells = []
    for prefix, label in CHAPTERS:
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=9, value=f'=SUMIF(A:A,"{prefix}.*",I:I)')
        for c in range(1, 10):
            apply_cell_style(ws.cell(row=r, column=c), "summary_chapter_row", c)
        chapter_total_cells.append(f"I{r}")
        r += 1
    total_row = r
    ws.cell(row=r, column=2, value='סה"כ עלות')
    ws.cell(row=r, column=9, value=f"=SUM({chapter_total_cells[0]}:{chapter_total_cells[-1]})")
    for c in range(1, 10):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)
    r += 1
    vat_row = r
    ws.cell(row=r, column=2, value='מע"מ בשיעור 18%')
    ws.cell(row=r, column=9, value=f"=(I{total_row}*(1/100*18))")
    for c in range(1, 10):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)
    r += 1
    ws.cell(row=r, column=2, value='סה"כ עלות כולל מע"מ')
    ws.cell(row=r, column=9, value=f"=I{total_row}+I{vat_row}")
    for c in range(1, 10):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)

    # ---- raw input block (columns K/L) ----
    ws["K1"] = "נתוני מדידה גולמיים (קלט מריכוז כמויות)"
    ws["K1"].font = Font(bold=True)
    ws["K2"] = "שדה"
    ws["L2"] = "ערך"
    ws["K2"].font = BOLD
    ws["L2"].font = BOLD
    field_row = {}
    rr = 3
    for field in ALL_RAW_FIELDS:
        ws.cell(row=rr, column=11, value=field)
        val_cell = ws.cell(row=rr, column=12, value=raw_values.get(field, 0))
        val_cell.font = Font(color="0000FF")
        field_row[field] = rr
        rr += 1

    def L(field):
        if field not in field_row:
            raise KeyError(f"unknown raw field: {field}")
        return f"L{field_row[field]}"

    def D(code):
        return f"D{code_to_row[code]}"

    def E(code):
        return f"E{code_to_row[code]}"

    def I(code):
        return f"I{code_to_row[code]}"

    always_show = set()
    for code, (kind, arg) in RULES.items():
        row = code_to_row[code]
        if kind == "raw":
            ws.cell(row=row, column=4, value=f"={L(arg)}")
        elif kind == "raw_sum":
            ws.cell(row=row, column=4, value="=" + "+".join(L(f) for f in arg))
        elif kind == "const":
            ws.cell(row=row, column=4, value=arg)
        elif kind == "formula":
            expr = arg

            def repl(m):
                kind2, key = m.group(1), m.group(2)
                if kind2 == "D":
                    return D(key)
                if kind2 == "E":
                    return E(key)
                if kind2 == "I":
                    return I(key)
                if kind2 == "L":
                    return L(key)
                raise ValueError(m.group(0))

            expr = re.sub(r"\{(D|E|I|L):([^}]+)\}", repl, expr)
            ws.cell(row=row, column=4, value="=" + expr)
        elif kind == "always0_visible":
            ws.cell(row=row, column=4, value=0)
            always_show.add(row)

    # ---- apply manual per-order corrections on top of the rule engine ----
    for code, override_value in overrides.items():
        if code not in code_to_row:
            raise KeyError(f"override for unknown section code: {code}")
        if isinstance(override_value, str) and "{" in override_value:
            expr = override_value.lstrip("=")

            def orepl(m):
                k, key = m.group(1), m.group(2)
                return {"D": D, "E": E, "I": I, "L": L}[k](key)

            override_value = "=" + re.sub(r"\{(D|E|I|L):([^}]+)\}", orepl, expr)
        ws.cell(row=code_to_row[code], column=4, value=override_value)

    return ws, code_to_row, last_item_row, always_show


def build_rikuz_sheet(wb, order_sheet_title, order_name, order_number, code_to_row, last_item_row):
    ws = wb.create_sheet("ריכוז")
    wb.move_sheet("ריכוז", offset=-(len(wb.sheetnames) - 1))  # put first
    ws.sheet_view.rightToLeft = True

    set_column_widths(ws, "ABCDEFG")
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 16
    ws.column_dimensions["K"].width = 16

    ws["A1"] = "חשבון מס'"
    ws["A1"].font = BOLD
    ws["B1"] = order_number
    ws["B1"].font = BOLD
    set_row_height(ws, 1)

    headers = ["סעיף", "תאור", "יח'", "מחיר", 'הנחת קבלן 15%', 'תוספת דירה מאוכלסת 5%', 'סה"כ מחיר יחידה']
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        apply_cell_style(cell, "header_row", c)

    ws.cell(row=3, column=8, value=order_name).font = BOLD
    ws.cell(row=3, column=9, value=order_number).font = BOLD
    ws.cell(row=4, column=8, value="כמות").font = BOLD
    ws.cell(row=4, column=8).fill = HEADER_FILL
    ws.cell(row=4, column=9, value='סה"כ לדירה').font = BOLD
    ws.cell(row=4, column=9).fill = HEADER_FILL
    ws.cell(row=4, column=10, value='סה"כ כמות חשבון').font = BOLD
    ws.cell(row=4, column=10).fill = HEADER_FILL
    ws.cell(row=4, column=11, value='סה"כ סכום חשבון').font = BOLD
    ws.cell(row=4, column=11).fill = HEADER_FILL
    set_row_height(ws, 4)

    for code, desc, unit, price in CANONICAL_ROWS:
        if code is None:
            continue
        r = code_to_row[code]
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=price if price != "" else None)
        if price != "":
            ws.cell(row=r, column=5, value=f"=IF(D{r}=\"\",\"\",D{r}*0.15)")
            ws.cell(row=r, column=6, value=f"=IF(D{r}=\"\",\"\",D{r}*0.05)")
            ws.cell(row=r, column=7, value=f"=IF(D{r}=\"\",\"\",D{r}*1.05*0.85)")
            ws.cell(row=r, column=8, value=f"='{order_sheet_title}'!D{r}")
            ws.cell(row=r, column=9, value=f"=IF(H{r}=\"\",\"\",H{r}*G{r})")
            ws.cell(row=r, column=10, value=f"=H{r}")
            ws.cell(row=r, column=11, value=f"=I{r}")
        for c in range(1, 8):
            apply_cell_style(ws.cell(row=r, column=c), "item_row", c)
        for c in (8, 9, 10, 11):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.number_format = "#,##0.00"
        set_row_height(ws, r)

    r = last_item_row + 2
    chapter_total_cells = []
    for prefix, label in CHAPTERS:
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=11, value=f'=SUMIF(A:A,"{prefix}.*",K:K)')
        for c in range(1, 12):
            apply_cell_style(ws.cell(row=r, column=c), "summary_chapter_row", c if c <= 9 else 9)
        chapter_total_cells.append(f"K{r}")
        r += 1
    total_row = r
    ws.cell(row=r, column=2, value='סה"כ עלות')
    ws.cell(row=r, column=11, value=f"=SUM({chapter_total_cells[0]}:{chapter_total_cells[-1]})")
    for c in range(1, 12):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c if c <= 9 else 9)
    r += 1
    vat_row = r
    ws.cell(row=r, column=2, value='מע"מ בשיעור 18%')
    ws.cell(row=r, column=11, value=f"=(K{total_row}*(1/100*18))")
    for c in range(1, 12):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c if c <= 9 else 9)
    r += 1
    ws.cell(row=r, column=2, value='סה"כ עלות כולל מע"מ')
    ws.cell(row=r, column=11, value=f"=K{total_row}+K{vat_row}")
    for c in range(1, 12):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c if c <= 9 else 9)

    return ws


def generate(order_name, order_number, raw_values, out_path, overrides=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_title = str(order_number)
    ws_order, code_to_row, last_row, always_show = build_order_sheet(
        wb, sheet_title, order_name, order_number, raw_values, overrides)
    build_rikuz_sheet(wb, sheet_title, order_name, order_number, code_to_row, last_row)
    wb.save(out_path)
    return code_to_row, always_show


if __name__ == "__main__":
    pass
