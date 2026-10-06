import copy
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from order_template import CANONICAL_ROWS
from generate import build_order_sheet, CHAPTERS, BOLD, HEADER_FILL, BOX
from styles import apply_cell_style, set_column_widths, set_row_height


def copy_full_sheet(src_ws, dst_wb, dst_title):
    """Verbatim copy of a worksheet (values/formulas, styles, merges, column
    widths, row heights, and hidden-row state) into a new sheet in dst_wb."""
    dst = dst_wb.create_sheet(dst_title)
    dst.sheet_view.rightToLeft = src_ws.sheet_view.rightToLeft

    for col, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst.column_dimensions[col].width = dim.width
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[r].height = dim.height
        if dim.hidden:
            dst.row_dimensions[r].hidden = True

    for merged_range in src_ws.merged_cells.ranges:
        dst.merge_cells(str(merged_range))

    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            nc = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            nc.font = copy.copy(cell.font)
            nc.border = copy.copy(cell.border)
            nc.alignment = copy.copy(cell.alignment)
            nc.number_format = cell.number_format
            if cell.fill:
                nc.fill = copy.copy(cell.fill)
    return dst


def build_address_list_sheet(wb, job):
    """New sheet: 'רשימת כתובות', placed right after 'ריכוז' and before the
    first order sheet. Columns: מס' הזמנה, שם משפחה, כתובת, קומה, דירה -
    one row per order, in job order. Pulled straight from job['meta'],
    nothing new to read."""
    ws = wb.create_sheet("רשימת כתובות")
    ws.sheet_view.rightToLeft = True

    HEADER_FONT = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    HDR_FILL = PatternFill("solid", fgColor="2E4A5E")
    BODY_FONT = Font(name="Arial", size=11)
    THIN = Side(style="thin", color="C9C9C9")
    CELL_BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    CENTER = Alignment(horizontal="center", vertical="center")
    RIGHT = Alignment(horizontal="right", vertical="center")

    headers = ["מס' הזמנה", "שם משפחה", "כתובת", "קומה", "דירה"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = CELL_BORDER
    ws.row_dimensions[1].height = 22

    r = 2
    for order in job["orders"]:
        meta = job["meta"].get(str(order["number"]), {})
        vals = [
            order["number"],
            order["name"],
            meta.get("כתובת") or "",
            meta.get("קומה") or "",
            meta.get("מספר דירה") or "",
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            cell.alignment = CENTER if c in (1, 4, 5) else RIGHT
        r += 1

    widths = [12, 20, 24, 8, 8]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    return ws


def build_combined_rikuz(wb, orders_meta, extra_rows_by_order=None):
    """orders_meta: list of (sheet_title, order_name, order_number, code_to_row)
    extra_rows_by_order: {sheet_title: [{'code','desc','unit','price','own_sheet_row'}]}
    - one-off rows that only exist in that specific order's own sheet, shown
    in the ריכוז but contributing only to that order's column pair."""
    extra_rows_by_order = extra_rows_by_order or {}

    ws = wb.create_sheet("ריכוז")
    wb.move_sheet("ריכוז", offset=-(len(wb.sheetnames) - 1))
    ws.sheet_view.rightToLeft = True

    set_column_widths(ws, "ABCDEFG")

    headers = ["סעיף", "תאור", "יח'", "מחיר", 'הנחת קבלן 15%', 'תוספת דירה מאוכלסת 5%', 'סה"כ מחיר יחידה']
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        apply_cell_style(cell, "header_row", c)

    ws["A1"] = "חשבון מרוכז"
    ws["A1"].font = BOLD
    set_row_height(ws, 1)

    code_to_row = orders_meta[0][3]

    col = 8
    order_col_pairs = []
    for sheet_title, order_name, order_number, _ in orders_meta:
        qty_col = col
        tot_col = col + 1
        ws.cell(row=3, column=qty_col, value=order_name).font = BOLD
        ws.cell(row=3, column=tot_col, value=order_number).font = BOLD
        ws.cell(row=4, column=qty_col, value="כמות").font = BOLD
        ws.cell(row=4, column=qty_col).fill = HEADER_FILL
        ws.cell(row=4, column=tot_col, value='סה"כ לדירה').font = BOLD
        ws.cell(row=4, column=tot_col).fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(qty_col)].width = 10
        ws.column_dimensions[get_column_letter(tot_col)].width = 14
        order_col_pairs.append((sheet_title, qty_col, tot_col))
        col += 2

    total_qty_col = col
    total_sum_col = col + 1
    ws.cell(row=4, column=total_qty_col, value='סה"כ כמות חשבון').font = BOLD
    ws.cell(row=4, column=total_qty_col).fill = HEADER_FILL
    ws.cell(row=4, column=total_sum_col, value='סה"כ סכום חשבון').font = BOLD
    ws.cell(row=4, column=total_sum_col).fill = HEADER_FILL
    ws.column_dimensions[get_column_letter(total_qty_col)].width = 16
    ws.column_dimensions[get_column_letter(total_sum_col)].width = 16
    set_row_height(ws, 4)

    total_sum_letter = get_column_letter(total_sum_col)

    for code, desc, unit, price in CANONICAL_ROWS:
        if code is None:
            continue
        r = code_to_row[code]
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=price if price != "" else None)
        for c in range(1, 8):
            apply_cell_style(ws.cell(row=r, column=c), "item_row", c)
        set_row_height(ws, r)

        if price == "":
            continue

        ws.cell(row=r, column=5, value=f"=IF(D{r}=\"\",\"\",D{r}*0.15)")
        ws.cell(row=r, column=6, value=f"=IF(D{r}=\"\",\"\",D{r}*0.05)")
        ws.cell(row=r, column=7, value=f"=IF(D{r}=\"\",\"\",D{r}*1.05*0.85)")

        qty_letters = []
        tot_letters = []
        for sheet_title, qty_col, tot_col in order_col_pairs:
            qty_letter = get_column_letter(qty_col)
            tot_letter = get_column_letter(tot_col)
            ws.cell(row=r, column=qty_col, value=f"='{sheet_title}'!D{r}")
            ws.cell(row=r, column=tot_col, value=f"=IF({qty_letter}{r}=\"\",\"\",{qty_letter}{r}*G{r})")
            for c in (qty_col, tot_col):
                cell = ws.cell(row=r, column=c)
                cell.border = BOX
                cell.number_format = "#,##0.00"
            qty_letters.append(f"{qty_letter}{r}")
            tot_letters.append(f"{tot_letter}{r}")

        ws.cell(row=r, column=total_qty_col, value="=" + "+".join(qty_letters))
        ws.cell(row=r, column=total_sum_col, value="=" + "+".join(tot_letters))
        for c in (total_qty_col, total_sum_col):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.number_format = "#,##0.00"

    last_item_row = max(code_to_row.values())
    r = last_item_row + 1

    for sheet_title, extras in extra_rows_by_order.items():
        for extra in extras:
            r += 1
            ws.cell(row=r, column=1, value=extra["code"])
            ws.cell(row=r, column=2, value=extra["desc"])
            ws.cell(row=r, column=3, value=extra["unit"])
            ws.cell(row=r, column=4, value=extra["price"])
            for c in range(1, 8):
                apply_cell_style(ws.cell(row=r, column=c), "item_row", c)
            ws.cell(row=r, column=5, value=f"=IF(D{r}=\"\",\"\",D{r}*0.15)")
            ws.cell(row=r, column=6, value=f"=IF(D{r}=\"\",\"\",D{r}*0.05)")
            ws.cell(row=r, column=7, value=f"=IF(D{r}=\"\",\"\",D{r}*1.05*0.85)")

            qty_letters = []
            tot_letters = []
            for st, qty_col, tot_col in order_col_pairs:
                qty_letter = get_column_letter(qty_col)
                tot_letter = get_column_letter(tot_col)
                if st == sheet_title:
                    ws.cell(row=r, column=qty_col,
                            value=f"='{st}'!D{extra['own_sheet_row']}")
                else:
                    ws.cell(row=r, column=qty_col, value=0)
                ws.cell(row=r, column=tot_col, value=f"=IF({qty_letter}{r}=\"\",\"\",{qty_letter}{r}*G{r})")
                for c in (qty_col, tot_col):
                    cell = ws.cell(row=r, column=c)
                    cell.border = BOX
                    cell.number_format = "#,##0.00"
                qty_letters.append(f"{qty_letter}{r}")
                tot_letters.append(f"{tot_letter}{r}")

            ws.cell(row=r, column=total_qty_col, value="=" + "+".join(qty_letters))
            ws.cell(row=r, column=total_sum_col, value="=" + "+".join(tot_letters))
            for c in (total_qty_col, total_sum_col):
                cell = ws.cell(row=r, column=c)
                cell.border = BOX
                cell.number_format = "#,##0.00"

    last_item_row = r

    # per-order chapter summary block
    r = last_item_row + 2
    per_order_summary_start = r
    for prefix, label in CHAPTERS:
        ws.cell(row=r, column=2, value=label)
        for c in range(1, 8):
            apply_cell_style(ws.cell(row=r, column=c), "summary_chapter_row", c)
        for sheet_title, qty_col, tot_col in order_col_pairs:
            tot_letter = get_column_letter(tot_col)
            cell = ws.cell(row=r, column=tot_col,
                            value=f'=SUMIF(A:A,"{prefix}.*",{tot_letter}:{tot_letter})')
            cell.border = BOX
            cell.number_format = "#,##0.00"
        r += 1
    per_order_total_row = r
    ws.cell(row=r, column=2, value='סה"כ עלות')
    for c in range(1, 8):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)
    for sheet_title, qty_col, tot_col in order_col_pairs:
        tot_letter = get_column_letter(tot_col)
        cell = ws.cell(row=r, column=tot_col,
                        value=f"=SUM({tot_letter}{per_order_summary_start}:{tot_letter}{r-1})")
        cell.border = BOX
        cell.number_format = "#,##0.00"
    r += 1
    per_order_vat_row = r
    ws.cell(row=r, column=2, value='מע"מ בשיעור 18%')
    for c in range(1, 8):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)
    for sheet_title, qty_col, tot_col in order_col_pairs:
        tot_letter = get_column_letter(tot_col)
        cell = ws.cell(row=r, column=tot_col,
                        value=f"=({tot_letter}{per_order_total_row}*(1/100*18))")
        cell.border = BOX
        cell.number_format = "#,##0.00"
    r += 1
    ws.cell(row=r, column=2, value='סה"כ עלות כולל מע"מ')
    for c in range(1, 8):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", c)
    for sheet_title, qty_col, tot_col in order_col_pairs:
        tot_letter = get_column_letter(tot_col)
        cell = ws.cell(row=r, column=tot_col,
                        value=f"={tot_letter}{per_order_total_row}+{tot_letter}{per_order_vat_row}")
        cell.border = BOX
        cell.number_format = "#,##0.00"

    last_item_row = r

    # whole-account chapter summary
    r = last_item_row + 2
    chapter_total_cells = []
    for prefix, label in CHAPTERS:
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=total_sum_col, value=f'=SUMIF(A:A,"{prefix}.*",{total_sum_letter}:{total_sum_letter})')
        for c in range(1, total_sum_col + 1):
            apply_cell_style(ws.cell(row=r, column=c), "summary_chapter_row", min(c, 9))
        chapter_total_cells.append(f"{total_sum_letter}{r}")
        r += 1
    total_row = r
    ws.cell(row=r, column=2, value='סה"כ עלות')
    ws.cell(row=r, column=total_sum_col, value=f"=SUM({chapter_total_cells[0]}:{chapter_total_cells[-1]})")
    for c in range(1, total_sum_col + 1):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", min(c, 9))
    r += 1
    vat_row = r
    ws.cell(row=r, column=2, value='מע"מ בשיעור 18%')
    ws.cell(row=r, column=total_sum_col, value=f"=({total_sum_letter}{total_row}*(1/100*18))")
    for c in range(1, total_sum_col + 1):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", min(c, 9))
    r += 1
    ws.cell(row=r, column=2, value='סה"כ עלות כולל מע"מ')
    ws.cell(row=r, column=total_sum_col, value=f"={total_sum_letter}{total_row}+{total_sum_letter}{vat_row}")
    for c in range(1, total_sum_col + 1):
        apply_cell_style(ws.cell(row=r, column=c), "summary_total_row", min(c, 9))

    return ws, code_to_row, last_item_row


def generate_multi(job, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    orders_meta = []
    extra_rows_by_order = {}
    for order in job["orders"]:
        sheet_title = str(order["number"])
        if order.get("locked"):
            src_wb = openpyxl.load_workbook(order["locked_file"], data_only=False)
            copy_full_sheet(src_wb[sheet_title], wb, sheet_title)
            from apply_hide import build_row_layout
            code_to_row, _ = build_row_layout()
            if order.get("extra_rows"):
                extra_rows_by_order[sheet_title] = order["extra_rows"]
        else:
            ws_order, code_to_row, last_row, always_show = build_order_sheet(
                wb, sheet_title, order["name"], order["number"], order["raw"],
                order.get("overrides"))
        orders_meta.append((sheet_title, order["name"], order["number"], code_to_row))

    build_combined_rikuz(wb, orders_meta, extra_rows_by_order)

    # NEW: address list sheet, right after ריכוז, before the first order sheet
    build_address_list_sheet(wb, job)
    # ריכוז is currently at position 0 (moved there in build_combined_rikuz);
    # move the address list sheet to position 1 (right after it).
    wb.move_sheet("רשימת כתובות", offset=-(len(wb.sheetnames) - 2))

    wb.save(out_path)
    return orders_meta
