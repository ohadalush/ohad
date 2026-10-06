"""
Per-element bracket assignment -> suggested per-order overrides.

rules.py maps each aggregate field to ONE default code (e.g. all sliding area
-> 12.011.0270, the 2-3 m² bracket). Real orders have elements in other
brackets, so every order needs overrides. This module reads 'פירוט אלמנטים',
assigns each element to its bracket by UNIT area + track count, and returns
only the overrides that differ from what rules.py would produce anyway.

It never guesses: anything ambiguous comes back as a warning with level
"ask" - stop and ask the user before building.
"""
import openpyxl

from order_template import CANONICAL_ROWS

PRICE = {c: p for c, d, u, p in CANONICAL_ROWS if c and p != ""}

# (code, lo, hi): element belongs to the first bracket with unit_area <= hi.
# Below the lowest -> lowest bracket; above the highest -> highest bracket.
# Quantity is always the REAL area (never clamped). Confirmed by the user.
SLIDING_2 = [("12.011.0260", 0.6, 1.0), ("12.011.0265", 1.0, 2.0), ("12.011.0270", 2.0, 3.0),
             ("12.011.0275", 3.0, 4.0), ("12.011.0280", 4.0, 5.0), ("12.011.0285", 5.0, 6.0)]
SLIDING_3 = [("12.012.0100", 1.0, 2.0), ("12.012.0110", 2.0, 3.0), ("12.012.0120", 3.0, 5.0),
             ("12.012.0130", 5.0, 7.0)]
VITRINA_3 = [("12.052.0070", 6.0, 8.0), ("12.052.0080", 8.0, 10.0)]
VITRINA_2 = [("12.051.0460", 4.0, 6.0), ("12.051.0470", 6.0, 8.0)]
SHUT_REG = [("12.065.1611", 0.6, 1.0), ("12.065.1613", 1.0, 3.0), ("12.065.1615", 3.0, 6.0)]
SHUT_LIGHT_WIN = [("12.065.1700", 0.6, 1.0), ("12.065.1710", 1.0, 3.0), ("12.065.1720", 3.0, 6.0),
                  ("12.065.1730", 6.0, 8.5)]
SHUT_LIGHT_DOOR = [("12.065.1661", 0.6, 1.0), ("12.065.1671", 1.0, 3.0), ("12.065.1681", 3.0, 6.0)]

FAMILY_CODES = sorted({c for fam in (SLIDING_2, SLIDING_3, VITRINA_3, VITRINA_2, SHUT_REG,
                                     SHUT_LIGHT_WIN, SHUT_LIGHT_DOOR) for c, _, _ in fam}
                      | {"12.011.1001", "12.014.0660"})

WINDOW_TYPES = {"הזזה", "הזזה+קבוע", "קיפ", "דרייקיפ"}
DOOR_TYPES = {"ויטרינה", "דלת ציר"}
BOX_ONLY_TYPES = {"טיפול בארגז"}
ASK_TYPES = {"קיפ+קבוע", "דרייקיפ+קבוע", "חלון ציר"}

# address -> (shutter type, rotating handles)
ADDRESSES = {
    "שולמית אלוני 15": ("אור", False),
    "משה רינת 4": ("אור", False),
    "משה רינת 10": ("אור", False),
    "מנחם בגין 40": ("אור", False),
    "מנחם בגין 50": ("אור", False),
    "מנחם בגין 52": ("אור", False),
    "משה רינת 2": ("רגיל", True),
    "מנחם בגין 48": ("רגיל", True),
}

GLASS_TARGET_WITH_BOX = ("=-{D:12.068.1172}-{D:12.068.1004}-{D:12.068.1006}"
                         "-{L:קיזוז זכוכית לארגז תריס}")


def pick(brackets, area):
    for code, lo, hi in brackets:
        if area <= hi + 1e-9:
            return code
    return brackets[-1][0]


def read_elements(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    if "פירוט אלמנטים" not in wb.sheetnames:
        return []
    ws = wb["פירוט אלמנטים"]
    header, out = None, []
    for row in ws.iter_rows(values_only=True):
        if not any(v is not None for v in row):
            continue
        if row[0] == "סוג אלמנט":
            header = [str(h).strip() if h else "" for h in row]
            continue
        if header is None:
            continue
        rec = dict(zip(header, row))
        t = str(rec.get("סוג אלמנט") or "").strip()
        if not t:
            continue

        def num(k):
            v = rec.get(k)
            try:
                return float(v)
            except (TypeError, ValueError):
                return 0.0
        out.append({
            "type": t,
            "count": num("כמות") or 1.0,
            "unit_area": num("שטח יחידה"),
            "shutter_unit_area": num("שטח יחידת תריס לפירוק"),
            "tracks": str(rec.get("מספר מסלולים") or "").strip(),
            "shutter": bool(str(rec.get("תריס") or "").strip()),
            "light": bool(str(rec.get("תריס אור") or "").strip()),
            "motor": rec.get("מנוע"),
            "notes": rec.get("הערות"),
        })
    return out


def _default_value(code, raw):
    """What rules.py would put in D for this code (only for the codes we manage)."""
    from rules import RULES
    kind, arg = RULES.get(code, ("const", 0))
    if kind == "raw":
        return raw.get(arg, 0) or 0
    if kind == "raw_sum":
        return sum(raw.get(f, 0) or 0 for f in arg)
    if kind == "const":
        return arg
    if code == "12.011.1001":
        return (raw.get("שטח הזזה+קבוע", 0) or 0) * PRICE["12.011.0270"] + \
               (raw.get("שטח קיפ+קבוע", 0) or 0) * PRICE["12.014.0660"]
    return None


def suggest(raw, meta, elements):
    warn = []          # (level, text)   level: "ask" = must stop, "info" = FYI
    target = {c: 0.0 for c in FAMILY_CODES}
    groups = {}        # summary field -> list of (code, area) so we can rescale

    def add(field, code, area):
        groups.setdefault(field, []).append([code, area])

    types = {e["type"] for e in elements}
    for t in types - WINDOW_TYPES - DOOR_TYPES - BOX_ONLY_TYPES:
        if t in ASK_TYPES:
            warn.append(("ask", f"סוג אלמנט '{t}' - צריך לשאול איך לפצל/לתמחר (לא אוטומטי)"))
        else:
            warn.append(("ask", f"סוג אלמנט לא מוכר: '{t}'"))

    sliding_fixed_base = 0.0
    kip_total = 0.0
    for e in elements:
        t, n, a, tr = e["type"], e["count"], e["unit_area"], e["tracks"]
        if t in ("הזזה", "הזזה+קבוע"):
            fam = SLIDING_3 if tr == "3" else SLIDING_2
            if tr not in ("2", "3"):
                warn.append(("ask", f"{t} {a} מ\"ר: מספר מסלולים לא ידוע ('{tr}')"))
            code = pick(fam, a)
            if code == "12.011.0280":
                warn.append(("ask", "12.011.0280 (זכוכית 6+10+10) - אין לו דלי זכוכית, לשאול לפני שימוש"))
            if a > fam[-1][2]:
                warn.append(("info", f"{t} {a} מ\"ר מעל הסוגר העליון -> {code} בשטח אמיתי"))
            field = "שטח הזזה" if t == "הזזה" else "שטח הזזה+קבוע"
            add(field, code, n * a)
        elif t == "ויטרינה":
            fam = VITRINA_2 if tr == "2" else VITRINA_3
            if tr not in ("2", "3"):
                warn.append(("ask", f"ויטרינה {a} מ\"ר: מספר מסלולים לא ידוע ('{tr}')"))
            code = pick(fam, a)
            if a > fam[-1][2] or a < fam[0][1]:
                warn.append(("info", f"ויטרינה {a} מ\"ר מחוץ לטווח הסוגריים -> {code} בשטח אמיתי"))
            add("שטח ויטרינה", code, n * a)
        elif t in ("קיפ", "דרייקיפ"):
            kip_total += n * max(0.6, a)
        elif t == "דלת ציר":
            if a > 3.0:
                warn.append(("ask", f"דלת ציר {a} מ\"ר - מעל 3 מ\"ר (חד/דו-כנפית?) לשאול"))

        # shutters
        if e["shutter"] or e["light"]:
            sa = a or e["shutter_unit_area"]
            door = t in DOOR_TYPES
            if e["light"]:
                fam = SHUT_LIGHT_DOOR if door else SHUT_LIGHT_WIN
                field = "שטח תריסים לדלתות (אור)" if door else "שטח תריסים לחלונות (אור)"
            else:
                fam = SHUT_REG
                field = "שטח תריסים לדלתות" if door else "שטח תריסים לחלונות"
            add(field, pick(fam, sa), n * sa)

    # rescale each group so it sums exactly to the summary field (element
    # sheet is rounded to 2 decimals; summary carries the exact value)
    for field, items in groups.items():
        el_total = sum(x[1] for x in items)
        summ = raw.get(field, 0) or 0
        if el_total and summ and abs(summ - el_total) / el_total < 0.02:
            k = summ / el_total
            for x in items:
                x[1] *= k
        elif el_total and abs(summ - el_total) > 0.05:
            warn.append(("ask", f"'{field}': סיכומים={summ} מול פירוט אלמנטים={round(el_total, 3)} - לא תואם"))
        for code, area in items:
            target[code] += area
        if field == "שטח הזזה+קבוע":
            sliding_fixed_base += sum(area * PRICE[code] for code, area in items)

    # summary fields that have area but no elements at all
    for field in ("שטח הזזה", "שטח הזזה+קבוע", "שטח ויטרינה", "שטח תריסים לחלונות",
                  "שטח תריסים לדלתות", "שטח תריסים לחלונות (אור)", "שטח תריסים לדלתות (אור)"):
        if (raw.get(field) or 0) and field not in groups:
            warn.append(("ask", f"'{field}'={raw[field]} בסיכומים אבל אין אלמנטים תואמים בפירוט"))

    if kip_total:
        target["12.014.0660"] = kip_total
    target["12.011.1001"] = sliding_fixed_base
    if (raw.get("שטח קיפ+קבוע") or 0):
        warn.append(("ask", "יש 'שטח קיפ+קבוע' - חלוקת קיפ/קבוע ידנית (12.014.0660 + 12.016.06xx)"))

    overrides = {}
    for code in FAMILY_CODES:
        dv = _default_value(code, raw)
        tv = round(target[code], 6)
        if dv is None or abs(tv - dv) > 1e-6:
            overrides[code] = tv

    # glass: always subtract the shutter-box glass from the target glass
    if raw.get("קיזוז זכוכית לארגז תריס"):
        overrides["12.068.1216"] = GLASS_TARGET_WITH_BOX

    # address rules: rotating handles + shutter-type sanity check
    addr = str(meta.get("כתובת") or "").strip()
    if addr:
        if addr not in ADDRESSES:
            warn.append(("ask", f"כתובת לא מוכרת: '{addr}' (סוג תריס / ידיות?)"))
        else:
            stype, handles = ADDRESSES[addr]
            if handles:
                win = sum(e["count"] for e in elements if e["type"] in WINDOW_TYPES | ASK_TYPES)
                vit = sum(e["count"] for e in elements if e["type"] == "ויטרינה")
                if win:
                    overrides["12.99.0002"] = win
                if vit:
                    overrides["12.99.0001"] = vit
                if any(e["type"] == "דלת ציר" for e in elements):
                    warn.append(("ask", "דלת ציר בכתובת עם ידיות סיבוביות - האם מקבלת ידית?"))
            has_light = any(e["light"] for e in elements)
            has_reg = any(e["shutter"] and not e["light"] for e in elements)
            if (stype == "אור" and has_reg) or (stype == "רגיל" and has_light):
                warn.append(("info", f"סוג התריס בפירוט לא תואם את ברירת המחדל של הכתובת ({stype}) - הולכים לפי הפירוט"))
    else:
        warn.append(("info", "אין כתובת בקובץ"))

    # special order: shutter-box treatment only
    area_fields = [k for k in raw if k.startswith("שטח") and (raw[k] or 0)]
    if not area_fields and (raw.get("כמות טיפול בארגז") or 0):
        overrides["60.020.0067"] = 1
        overrides["60.010.0020"] = 0
        warn.append(("info", "הזמנת 'ארגזי תריס בלבד': צוות אלומיניום יום 1, פועל פשוט 0"))

    if not elements:
        warn.append(("ask", "אין גיליון 'פירוט אלמנטים' / אין אלמנטים - אי אפשר לשייך סוגריים"))

    return overrides, warn
